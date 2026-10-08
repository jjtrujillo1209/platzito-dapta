"""Contactos: CRUD, importación CSV/XLSX con mapeo de columnas, línea de tiempo y opt-out."""
import csv
import io

import phonenumbers
from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from pydantic import BaseModel
from sqlalchemy import func, or_, select

from .. import serial
from ..eventos import emitir
from ..modelos import Cita, Contacto, Conversacion, Inscripcion, Llamada, Secuencia
from ..secuencias import al_dar_baja
from ..seguridad import Contexto, requiere

r = APIRouter(prefix="/api/contactos", tags=["contactos"])

# Prefijo → zona horaria por defecto para LatAm (cuando el contacto no trae zona)
ZONA_POR_PAIS = {"CO": "America/Bogota", "MX": "America/Mexico_City", "PE": "America/Lima",
                 "CL": "America/Santiago", "AR": "America/Argentina/Buenos_Aires", "EC": "America/Guayaquil",
                 "US": "America/New_York", "ES": "Europe/Madrid", "GT": "America/Guatemala",
                 "CR": "America/Costa_Rica", "PA": "America/Panama", "UY": "America/Montevideo",
                 "BO": "America/La_Paz", "PY": "America/Asuncion", "VE": "America/Caracas",
                 "DO": "America/Santo_Domingo", "SV": "America/El_Salvador", "HN": "America/Tegucigalpa"}


def normalizar_telefono(valor: str | None, region: str = "CO") -> str | None:
    if not valor:
        return None
    texto = str(valor).strip()
    if texto.endswith(".0"):  # números que vienen de Excel como float
        texto = texto[:-2]
    try:
        num = phonenumbers.parse(texto if texto.startswith("+") else texto, region)
        if not phonenumbers.is_possible_number(num):
            if not texto.startswith("+"):
                num = phonenumbers.parse("+" + texto, None)
            if not phonenumbers.is_possible_number(num):
                return None
        return phonenumbers.format_number(num, phonenumbers.PhoneNumberFormat.E164)
    except phonenumbers.NumberParseException:
        return None


def zona_de_telefono(e164: str | None) -> str | None:
    if not e164:
        return None
    try:
        return ZONA_POR_PAIS.get(phonenumbers.region_code_for_number(phonenumbers.parse(e164)))
    except phonenumbers.NumberParseException:
        return None


def _region(ctx: Contexto) -> str:
    inv = {v: k for k, v in ZONA_POR_PAIS.items()}
    return inv.get(ctx.espacio.zona_horaria, "CO")


def _contacto(ctx: Contexto, contacto_id: int) -> Contacto:
    c = ctx.db.get(Contacto, contacto_id)
    if not c or c.espacio_id != ctx.espacio_id:
        raise HTTPException(404, "Contacto no encontrado")
    return c


@r.get("")
def listar(q: str = "", etapa: str = "", etiqueta: str = "", pagina: int = 1, por_pagina: int = 50,
           ctx: Contexto = Depends(requiere("lector"))):
    consulta = select(Contacto).where(Contacto.espacio_id == ctx.espacio_id)
    if q:
        like = f"%{q.strip()}%"
        consulta = consulta.where(or_(Contacto.nombre.ilike(like), Contacto.telefono.ilike(like),
                                      Contacto.email.ilike(like), Contacto.empresa.ilike(like)))
    if etapa:
        consulta = consulta.where(Contacto.etapa == etapa)
    filas = ctx.db.scalars(consulta.order_by(Contacto.id.desc())).all() if etiqueta else None
    if etiqueta:
        filas = [c for c in filas if etiqueta in (c.etiquetas or [])]
        total = len(filas)
        filas = filas[(pagina - 1) * por_pagina: pagina * por_pagina]
    else:
        total = ctx.db.scalar(select(func.count()).select_from(consulta.subquery()))
        filas = ctx.db.scalars(consulta.order_by(Contacto.id.desc()).offset((pagina - 1) * por_pagina)
                               .limit(min(por_pagina, 500))).all()
    etapas = [e for (e,) in ctx.db.execute(select(Contacto.etapa).where(Contacto.espacio_id == ctx.espacio_id)
                                           .distinct())]
    etiquetas = sorted({e for (lista,) in ctx.db.execute(select(Contacto.etiquetas).where(
        Contacto.espacio_id == ctx.espacio_id)) for e in (lista or [])})
    return {"total": total, "pagina": pagina, "contactos": [serial.contacto(c) for c in filas], "etapas": etapas,
            "etiquetas": etiquetas}


class DatosContacto(BaseModel):
    nombre: str | None = None
    telefono: str | None = None
    email: str | None = None
    empresa: str | None = None
    zona_horaria: str | None = None
    etapa: str | None = None
    atributos: dict | None = None
    etiquetas: list[str] | None = None
    opt_out_whatsapp: bool | None = None
    opt_out_llamadas: bool | None = None
    opt_out_correo: bool | None = None


def _aplicar(ctx: Contexto, c: Contacto, d: DatosContacto):
    datos = d.model_dump(exclude_none=True)
    if "telefono" in datos:
        tel = normalizar_telefono(datos.pop("telefono"), _region(ctx))
        if d.telefono and not tel:
            raise HTTPException(422, "Teléfono inválido")
        otro = ctx.db.scalar(select(Contacto).where(Contacto.espacio_id == ctx.espacio_id, Contacto.telefono == tel,
                                                    Contacto.id != (c.id or 0))) if tel else None
        if otro:
            raise HTTPException(409, f"Ya existe un contacto con ese teléfono (#{otro.id})")
        c.telefono = tel
        if not c.zona_horaria:
            c.zona_horaria = zona_de_telefono(tel)
    if "email" in datos:
        datos["email"] = datos["email"].strip().lower() or None
    for k, v in datos.items():
        setattr(c, k, v)


@r.post("")
def crear(datos: DatosContacto, ctx: Contexto = Depends(requiere("operador"))):
    c = Contacto(espacio_id=ctx.espacio_id)
    _aplicar(ctx, c, datos)
    if not (c.telefono or c.email):
        raise HTTPException(422, "El contacto necesita teléfono o correo")
    ctx.db.add(c)
    ctx.db.commit()
    return serial.contacto(c)


@r.get("/{contacto_id}")
def ver(contacto_id: int, ctx: Contexto = Depends(requiere("lector"))):
    c = _contacto(ctx, contacto_id)
    convs = ctx.db.scalars(select(Conversacion).where(Conversacion.contacto_id == c.id)
                           .order_by(Conversacion.id.desc()).limit(50)).all()
    llamadas = ctx.db.scalars(select(Llamada).where(Llamada.contacto_id == c.id).order_by(Llamada.id.desc())
                              .limit(50)).all()
    inscripciones = ctx.db.execute(select(Inscripcion, Secuencia.nombre).join(Secuencia)
                                   .where(Inscripcion.contacto_id == c.id)).all()
    citas = ctx.db.scalars(select(Cita).where(Cita.contacto_id == c.id).order_by(Cita.inicio.desc())).all()
    return {**serial.contacto(c), "conversaciones": [serial.conversacion(x) for x in convs],
            "llamadas": [serial.llamada(x) for x in llamadas],
            "inscripciones": [{**serial.inscripcion(i), "secuencia": n} for i, n in inscripciones],
            "citas": [serial.cita(x) for x in citas]}


@r.patch("/{contacto_id}")
def editar(contacto_id: int, datos: DatosContacto, ctx: Contexto = Depends(requiere("operador"))):
    c = _contacto(ctx, contacto_id)
    _aplicar(ctx, c, datos)
    ctx.db.commit()
    emitir(ctx.db, ctx.espacio_id, "contacto.actualizado", serial.contacto(c))
    return serial.contacto(c)


@r.delete("/{contacto_id}")
def borrar(contacto_id: int, ctx: Contexto = Depends(requiere("admin"))):
    c = _contacto(ctx, contacto_id)
    ctx.db.delete(c)
    ctx.db.commit()
    return {"ok": True}


@r.post("/{contacto_id}/baja")
def baja(contacto_id: int, datos: dict, ctx: Contexto = Depends(requiere("operador"))):
    c = _contacto(ctx, contacto_id)
    canales = datos.get("canales") or ["whatsapp", "llamadas", "correo"]
    for canal in canales:
        setattr(c, f"opt_out_{canal}", True)
    ctx.db.commit()
    if set(canales) >= {"whatsapp", "llamadas"}:
        al_dar_baja(ctx.db, c.id)
    emitir(ctx.db, ctx.espacio_id, "contacto.baja", {"contacto_id": c.id, "canales": canales})
    return serial.contacto(c)


# ─────────────────────────────── Importación ───────────────────────────────

ALIAS = {
    "nombre": ["nombre", "name", "nombre completo", "full name", "first name", "primer nombre"],
    "telefono": ["telefono", "teléfono", "phone", "celular", "móvil", "movil", "whatsapp", "phone number", "numero"],
    "email": ["email", "correo", "e-mail", "mail", "correo electrónico"],
    "empresa": ["empresa", "company", "compañía", "organizacion", "organización"],
    "zona_horaria": ["zona_horaria", "timezone", "zona horaria"],
    "etapa": ["etapa", "stage", "estado"],
}


def _leer_tabla(nombre: str, contenido: bytes) -> list[dict]:
    if nombre.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        hoja = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True).worksheets[0]
        filas = list(hoja.iter_rows(values_only=True))
        if not filas:
            return []
        cab = [str(c or "").strip() for c in filas[0]]
        return [{cab[i]: ("" if v is None else str(v)) for i, v in enumerate(f) if i < len(cab)} for f in filas[1:]]
    texto = contenido.decode("utf-8-sig", "ignore")
    dialecto = csv.Sniffer().sniff(texto[:4000], delimiters=",;\t") if texto.strip() else csv.excel
    return list(csv.DictReader(io.StringIO(texto), dialect=dialecto))


def mapeo_sugerido(columnas: list[str]) -> dict:
    mapeo = {}
    for col in columnas:
        clave = col.strip().lower()
        destino = next((campo for campo, alias in ALIAS.items() if clave in alias), None)
        mapeo[col] = destino or f"atributo:{clave.replace(' ', '_')}"
    return mapeo


@r.post("/importar/previsualizar")
async def previsualizar(archivo: UploadFile = File(...), ctx: Contexto = Depends(requiere("operador"))):
    filas = _leer_tabla(archivo.filename or "", await archivo.read())
    columnas = list(filas[0].keys()) if filas else []
    return {"columnas": columnas, "mapeo": mapeo_sugerido(columnas), "muestra": filas[:5], "total": len(filas)}


@r.post("/importar")
async def importar(archivo: UploadFile = File(...), mapeo: str = Form("{}"), etiquetas: str = Form(""),
                   secuencia_id: int | None = Form(None), ctx: Contexto = Depends(requiere("operador"))):
    import json

    filas = _leer_tabla(archivo.filename or "", await archivo.read())
    if not filas:
        raise HTTPException(422, "El archivo está vacío")
    m = json.loads(mapeo) or mapeo_sugerido(list(filas[0].keys()))
    tags = [t.strip() for t in etiquetas.split(",") if t.strip()]
    region = _region(ctx)
    creados = actualizados = invalidos = 0
    ids: list[int] = []
    existentes = {c.telefono: c for c in ctx.db.scalars(select(Contacto).where(
        Contacto.espacio_id == ctx.espacio_id, Contacto.telefono.is_not(None)))}
    por_email = {c.email: c for c in existentes.values() if c.email}
    for fila in filas:
        datos, atributos = {}, {}
        for col, destino in m.items():
            valor = (fila.get(col) or "").strip()
            if not valor or not destino or destino == "ignorar":
                continue
            if destino.startswith("atributo:"):
                atributos[destino[9:]] = valor
            else:
                datos[destino] = valor
        tel = normalizar_telefono(datos.get("telefono"), region)
        email = (datos.get("email") or "").lower() or None
        if not tel and not email:
            invalidos += 1
            continue
        c = existentes.get(tel) if tel else None
        c = c or (por_email.get(email) if email else None)
        if c:
            actualizados += 1
        else:
            c = Contacto(espacio_id=ctx.espacio_id, telefono=tel, zona_horaria=zona_de_telefono(tel))
            ctx.db.add(c)
            creados += 1
            if tel:
                existentes[tel] = c
        if email:
            c.email = email
            por_email[email] = c
        for campo in ("nombre", "empresa", "zona_horaria", "etapa"):
            if datos.get(campo):
                setattr(c, campo, datos[campo])
        c.atributos = {**(c.atributos or {}), **atributos}
        c.etiquetas = sorted(set((c.etiquetas or []) + tags))
        ctx.db.flush()
        ids.append(c.id)
    ctx.db.commit()
    inscritos = 0
    if secuencia_id:
        from ..secuencias import inscribir

        sec = ctx.db.get(Secuencia, secuencia_id)
        if sec and sec.espacio_id == ctx.espacio_id:
            inscritos = inscribir(ctx.db, sec, ids)
    return {"creados": creados, "actualizados": actualizados, "invalidos": invalidos, "inscritos": inscritos}
