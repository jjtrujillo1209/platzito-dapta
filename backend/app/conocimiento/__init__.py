"""Knowledge base ("cerebros"): ingesta → fragmentos → embeddings → búsqueda semántica.

Los embeddings se guardan como float32 normalizados y la búsqueda es coseno en
numpy, cacheada por cerebro. Funciona igual en SQLite y Postgres y alcanza
para decenas de miles de fragmentos; más allá conviene pgvector.
"""
import asyncio
import io
import re
import threading
from urllib.parse import urljoin, urlparse

import httpx
import numpy as np
from bs4 import BeautifulSoup
from defusedxml import ElementTree
from sqlalchemy import delete, func, select
from sqlalchemy.orm import Session

from ..config import ajustes
from ..db import ahora
from ..integraciones import credenciales
from ..modelos import Cerebro, Fragmento, Fuente
from ..red import validar_url_externa

TAM_FRAGMENTO = 3200  # ~800 tokens
SOLAPE = 400
CABECERAS_HTTP = {"user-agent": "Mozilla/5.0 (compatible; PlatzitoBot/1.0; +https://platzi.com)"}


class ErrorEmbeddings(Exception):
    pass


# ─────────────────────────────── Embeddings ───────────────────────────────

async def embeber(db: Session | None, espacio_id: int | None, textos: list[str], es_consulta: bool = False
                  ) -> tuple[np.ndarray, str]:
    """Devuelve matriz (n, d) normalizada y el nombre del modelo."""
    a = ajustes()
    prov, modelo = a.embeddings_proveedor, a.embeddings_modelo
    if not textos:
        return np.zeros((0, 1), dtype=np.float32), modelo
    vectores: list[list[float]] = []
    async with httpx.AsyncClient(timeout=120) as cliente:
        if prov == "ollama":
            url = credenciales(db, espacio_id, "ollama").get("url") or a.ollama_url
            if url.rstrip("/") != a.ollama_url.rstrip("/"):
                try:
                    validar_url_externa(url)  # la URL del espacio no es de confianza; la de la instancia sí
                except ValueError as e:
                    raise ErrorEmbeddings(str(e)) from e
            prefijo = "search_query: " if es_consulta else "search_document: "
            entradas = [prefijo + t if "nomic" in modelo else t for t in textos]
            for i in range(0, len(entradas), 32):
                try:
                    r = await cliente.post(f"{url}/api/embed", json={"model": modelo, "input": entradas[i:i + 32]})
                except httpx.HTTPError as e:
                    raise ErrorEmbeddings(f"No se pudo contactar Ollama en {url}: {e}") from e
                if r.status_code >= 400:
                    raise ErrorEmbeddings(f"Ollama {r.status_code}: {r.text[:200]}")
                vectores.extend(r.json()["embeddings"])
        elif prov == "openai":
            key = credenciales(db, espacio_id, "openai").get("api_key")
            if not key:
                raise ErrorEmbeddings("Falta la API key de OpenAI para embeddings")
            for i in range(0, len(textos), 96):
                r = await cliente.post("https://api.openai.com/v1/embeddings",
                                       headers={"authorization": f"Bearer {key}"},
                                       json={"model": modelo, "input": textos[i:i + 96]})
                if r.status_code >= 400:
                    raise ErrorEmbeddings(f"OpenAI {r.status_code}: {r.text[:200]}")
                vectores.extend(d["embedding"] for d in r.json()["data"])
        else:
            raise ErrorEmbeddings(f"Proveedor de embeddings desconocido: {prov}")
    m = np.asarray(vectores, dtype=np.float32)
    normas = np.linalg.norm(m, axis=1, keepdims=True)
    normas[normas == 0] = 1
    return m / normas, modelo


# ─────────────────────────────── Extracción de texto ───────────────────────────────

def html_a_texto(html: str) -> tuple[str, str]:
    sopa = BeautifulSoup(html, "html.parser")
    titulo = (sopa.title.string or "").strip() if sopa.title and sopa.title.string else ""
    for basura in sopa(["script", "style", "noscript", "svg", "iframe", "form", "nav", "footer", "header"]):
        basura.decompose()
    principal = sopa.find("main") or sopa.find("article") or sopa.body or sopa
    lineas = []
    for el in principal.find_all(["h1", "h2", "h3", "h4", "p", "li", "td", "th", "blockquote", "pre"]):
        txt = " ".join(el.get_text(" ", strip=True).split())
        if not txt:
            continue
        if el.name in ("h1", "h2", "h3", "h4"):
            txt = "#" * int(el.name[1]) + " " + txt
        elif el.name == "li":
            txt = "- " + txt
        lineas.append(txt)
    texto = "\n".join(lineas) if lineas else " ".join(principal.get_text(" ", strip=True).split())
    return titulo, texto


def archivo_a_texto(nombre: str, contenido: bytes) -> str:
    ext = nombre.lower().rsplit(".", 1)[-1]
    if ext == "pdf":
        from pypdf import PdfReader

        lector = PdfReader(io.BytesIO(contenido))
        return "\n\n".join((p.extract_text() or "") for p in lector.pages)
    if ext == "docx":
        import docx

        d = docx.Document(io.BytesIO(contenido))
        partes = [p.text for p in d.paragraphs if p.text.strip()]
        for tabla in d.tables:
            for fila in tabla.rows:
                partes.append(" | ".join(c.text.strip() for c in fila.cells))
        return "\n".join(partes)
    if ext in ("xlsx", "xlsm"):
        from openpyxl import load_workbook

        libro = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
        partes = []
        for hoja in libro.worksheets:
            partes.append(f"# {hoja.title}")
            for fila in hoja.iter_rows(values_only=True):
                valores = [str(v) for v in fila if v is not None]
                if valores:
                    partes.append(" | ".join(valores))
        return "\n".join(partes)
    if ext in ("html", "htm"):
        return html_a_texto(contenido.decode("utf-8", "ignore"))[1]
    if ext in ("txt", "md", "csv", "json", "xml"):
        return contenido.decode("utf-8", "ignore")
    raise ValueError(f"Formato no soportado: .{ext} (usa pdf, docx, xlsx, csv, txt, md, html)")


def fragmentar(texto: str, tam: int = TAM_FRAGMENTO, solape: int = SOLAPE) -> list[str]:
    texto = re.sub(r"\n{3,}", "\n\n", texto).strip()
    if len(texto) <= tam:
        return [texto] if texto else []
    parrafos = re.split(r"\n(?=#)|\n\n", texto)
    fragmentos, actual = [], ""
    for p in parrafos:
        while len(p) > tam:  # párrafo gigante: cortar por frases
            corte = p.rfind(". ", 0, tam)
            corte = corte + 1 if corte > tam // 2 else tam
            p_ini, p = p[:corte], p[corte:]
            if actual:
                fragmentos.append(actual)
                actual = ""
            fragmentos.append(p_ini.strip())
        if len(actual) + len(p) + 2 > tam and actual:
            fragmentos.append(actual)
            actual = actual[-solape:] + "\n\n" + p if solape else p
        else:
            actual = f"{actual}\n\n{p}" if actual else p
    if actual.strip():
        fragmentos.append(actual)
    return [f.strip() for f in fragmentos if f.strip()]


# ─────────────────────────────── Rastreo web ───────────────────────────────

MAX_REDIRECCIONES = 5


async def obtener_seguro(cliente: httpx.AsyncClient, url: str) -> httpx.Response:
    """GET que sigue redirecciones a mano, validando cada salto contra SSRF."""
    for _ in range(MAX_REDIRECCIONES + 1):
        r = await cliente.get(validar_url_externa(url), headers=CABECERAS_HTTP, follow_redirects=False)
        if r.is_redirect and r.headers.get("location"):
            url = urljoin(str(r.url), r.headers["location"])
            continue
        return r
    raise ValueError(f"Demasiadas redirecciones: {url}")


async def descargar_pagina(cliente: httpx.AsyncClient, url: str) -> tuple[str, str]:
    r = await obtener_seguro(cliente, url)
    r.raise_for_status()
    return html_a_texto(r.text)


async def descubrir_sitemap(url: str, max_urls: int = 200) -> list[str]:
    """Acepta la URL del sitio o del sitemap; sigue índices de sitemaps."""
    base = url if url.endswith(".xml") else urljoin(url if url.endswith("/") else url + "/", "sitemap.xml")
    vistos, pendientes, urls = set(), [base], []
    async with httpx.AsyncClient(timeout=20, headers=CABECERAS_HTTP) as cliente:
        while pendientes and len(urls) < max_urls:
            actual = pendientes.pop(0)
            if actual in vistos:
                continue
            vistos.add(actual)
            try:
                r = await obtener_seguro(cliente, actual)
                raiz = ElementTree.fromstring(r.content)
            except Exception:
                continue
            ns = raiz.tag.split("}")[0] + "}" if raiz.tag.startswith("{") else ""
            if raiz.tag.endswith("sitemapindex"):
                pendientes += [e.text.strip() for e in raiz.iter(f"{ns}loc") if e.text]
            else:
                urls += [e.text.strip() for e in raiz.iter(f"{ns}loc") if e.text]
    dominio = urlparse(url).netloc
    return [u for u in dict.fromkeys(urls) if urlparse(u).netloc.endswith(dominio.removeprefix("www."))][:max_urls]


# ─────────────────────────────── Pipeline por fuente ───────────────────────────────

async def procesar_fuente(db: Session, fuente_id: int):
    fuente = db.get(Fuente, fuente_id)
    if not fuente:
        return
    cerebro = db.get(Cerebro, fuente.cerebro_id)
    fuente.estado, fuente.error = "procesando", ""
    db.commit()
    try:
        documentos: list[tuple[str, str, str]] = []  # (titulo, url, texto)
        if fuente.tipo == "texto":
            documentos.append((fuente.nombre, "", fuente.texto))
        elif fuente.tipo == "archivo":
            ruta = ajustes().directorio_datos / fuente.archivo
            documentos.append((fuente.nombre, "", archivo_a_texto(fuente.nombre, ruta.read_bytes())))
        elif fuente.tipo == "url":
            async with httpx.AsyncClient(timeout=30) as cliente:
                titulo, texto = await descargar_pagina(cliente, fuente.url)
            documentos.append((titulo or fuente.url, fuente.url, texto))
        elif fuente.tipo == "sitemap":
            urls = fuente.opciones.get("urls") or await descubrir_sitemap(
                fuente.url, int(fuente.opciones.get("max_paginas", 50)))
            if not urls:
                raise ValueError("No se encontraron páginas en el sitemap")
            sem = asyncio.Semaphore(5)

            async with httpx.AsyncClient(timeout=30) as cliente:
                async def una(u):
                    async with sem:
                        try:
                            t, txt = await descargar_pagina(cliente, u)
                            return (t or u, u, txt)
                        except Exception:
                            return None

                documentos = [d for d in await asyncio.gather(*(una(u) for u in urls)) if d]
        fragmentos = [(t, u, f) for t, u, texto in documentos for f in fragmentar(texto)]
        if not fragmentos:
            raise ValueError("La fuente no tiene texto utilizable")
        matriz, modelo = await embeber(db, cerebro.espacio_id, [f"{t}\n{f}" for t, _, f in fragmentos])
        db.execute(delete(Fragmento).where(Fragmento.fuente_id == fuente.id))
        for i, ((titulo, url, texto), vec) in enumerate(zip(fragmentos, matriz)):
            db.add(Fragmento(cerebro_id=cerebro.id, fuente_id=fuente.id, orden=i, texto=texto,
                             titulo=titulo[:300], url=url[:1000], embedding=vec.tobytes(), modelo=modelo))
        fuente.n_fragmentos = len(fragmentos)
        if fuente.tipo == "sitemap":
            fuente.texto = "\n".join(u for _, u, _ in documentos)[:20000]
        fuente.estado = "lista"
    except Exception as e:  # el error queda visible en la UI y se puede reintentar
        db.rollback()
        fuente = db.get(Fuente, fuente_id)
        fuente.estado, fuente.error = "fallida", str(e)[:500]
    fuente.actualizado = ahora()
    db.commit()
    _invalidar(fuente.cerebro_id)


# ─────────────────────────────── Búsqueda ───────────────────────────────

_cache: dict[int, tuple[tuple, np.ndarray, list[int]]] = {}
_cache_lock = threading.Lock()


def _invalidar(cerebro_id: int):
    with _cache_lock:
        _cache.pop(cerebro_id, None)


def _matriz(db: Session, cerebro_id: int) -> tuple[np.ndarray, list[int]]:
    firma = tuple(db.execute(select(func.count(Fragmento.id), func.max(Fragmento.id))
                             .where(Fragmento.cerebro_id == cerebro_id)).one())
    with _cache_lock:
        if cerebro_id in _cache and _cache[cerebro_id][0] == firma:
            return _cache[cerebro_id][1], _cache[cerebro_id][2]
    filas = db.execute(select(Fragmento.id, Fragmento.embedding).where(Fragmento.cerebro_id == cerebro_id)).all()
    if not filas:
        m, ids = np.zeros((0, 1), dtype=np.float32), []
    else:
        m = np.vstack([np.frombuffer(e, dtype=np.float32) for _, e in filas])
        ids = [i for i, _ in filas]
    with _cache_lock:
        _cache[cerebro_id] = (firma, m, ids)
    return m, ids


async def buscar(db: Session, espacio_id: int, cerebro_ids: list[int], consulta: str, top_k: int = 4,
                 umbral: float = 0.35) -> list[dict]:
    if not cerebro_ids or not consulta.strip():
        return []
    cerebro_ids = list(db.scalars(select(Cerebro.id).where(Cerebro.id.in_(cerebro_ids),
                                                           Cerebro.espacio_id == espacio_id)))
    if not cerebro_ids:
        return []
    q, _ = await embeber(db, espacio_id, [consulta], es_consulta=True)
    candidatos: list[tuple[float, int]] = []
    for cid in cerebro_ids:
        m, ids = _matriz(db, cid)
        if not ids or m.shape[1] != q.shape[1]:
            continue
        puntajes = m @ q[0]
        mejores = np.argsort(-puntajes)[: top_k * 2]
        candidatos += [(float(puntajes[i]), ids[i]) for i in mejores]
    candidatos = sorted(c for c in candidatos if c[0] >= umbral)[::-1][:top_k]
    if not candidatos:
        return []
    frags = {f.id: f for f in db.scalars(select(Fragmento).where(Fragmento.id.in_([i for _, i in candidatos])))}
    return [{"texto": frags[i].texto, "titulo": frags[i].titulo, "url": frags[i].url, "puntaje": round(p, 3)}
            for p, i in candidatos if i in frags]
